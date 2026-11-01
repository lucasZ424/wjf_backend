"""点表 Excel 与部署配置 JSON 的互转（部署时用；后端运行时不引用本工具）。

    python pointsheet.py template --base production.base.json --out 点表.xlsx
    python pointsheet.py convert  --base production.base.json --sheet 点表.xlsx --out production.json

点表只填现场事实（修复轮16，用户 2026-11-01）：每行一个订阅点，填“类别”、内外壁温度的“温度对号”、显示名称（位号）、
数据源与地址（Modbus“地址”或 OPC UA“NodeId”）；单位、阈值、极值可以不填。内部编号、派生量、界面测站都由本工具生成：

  * 内壁温度、外壁温度 → temp_inner_NN、temp_outer_NN（NN 为温度对号）；同一温度对号的一内一外是一对，生成内外壁温差
    delta_t_NN 与热应力估算 stress_NN；全部温度对再生成一个热应力代理绝对峰值 thermal_stress_peak。
  * 电桥 → bridge_mvv_NN（按表中出现的顺序编号），生成换算应变 strain_NN、应变变化量 strain_change_NN 和一个测站色块。
    电桥不与温度关联：页面上的测站只显示电桥这一路，一路电桥周围常不止一对温度。
  * 八类过程量各最多一个，编号固定（负荷 load_mw 等）；“其他”→ other_NN，照常接入，不生成派生量。
  * 单位、阈值、极值留空时按类别取默认值，转换时列出；要去掉某个默认阈值，在那一格填“无”。“其他”类没有默认值，单位必填。
  * “地址”或“NodeId”留空的行不接入，逐条列出；温度对只接入了一侧时照常接入那一侧，不生成温差与热应力，列出。
  * 有错（类别不认识、温度对号缺失或重复、过程量重复……）按“表名 第几行 哪一列”报错，不生成文件。最后按后端自己的配置
    规则整份校验（与 `htsg check` 相同），通过才写出文件。

template 生成空白点表：数据源表按基础配置填好编号，两张点表只有表头、下拉选项和填写说明，不预填任何点。基础配置只提供点表
以外的部分（项目、模型、时限、存储、界面精度、模拟设置……），它自带的点、派生量、测站和标签都不用。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pydantic import ValidationError

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from htsg.config import Point, Settings, Source, check_node_id  # noqa: E402

YES, NO, NONE = "是", "否", "无"
EXPLICIT_NONE = object()  # 单元格里写了“无”：不要默认值

# (表头, 配置字段, 类型)；类型：text / int / float / bool / floats（分号分隔的多个数）
SOURCE_COLUMNS = [
    ("编号", "id", "text"), ("类型", "kind", "text"), ("启用", "enabled", "bool"), ("主机", "host", "text"),
    ("端口", "port", "int"), ("端点", "endpoint", "text"), ("用户名环境变量", "username_env", "text"),
    ("口令环境变量", "password_env", "text"), ("安全策略", "security", "text"),
]  # fmt: skip
COMMON_HEAD = [("类别", "category", "text"), ("温度对号", "pair", "int"), ("显示名称", None, "text"),
               ("备注", None, "text"), ("数据源", "source", "text"), ("单位", "unit", "text")]  # fmt: skip
LIMIT_COLUMNS = [
    ("比例", "scale", "float"), ("偏移", "offset", "float"), ("无效值", "invalid_values", "floats"),
    ("下限", "low", "float"), ("下限回差", "hysteresis_low", "float"), ("上限", "high", "float"),
    ("上限回差", "hysteresis_high", "float"), ("极小值", "extreme_low", "float"), ("极大值", "extreme_high", "float"),
]  # fmt: skip
MODBUS_COLUMNS = [
    *COMMON_HEAD, ("从站号", "device_id", "int"), ("寄存器类型", "register_kind", "text"), ("地址", "address", "int"),
    ("数据类型", "data_type", "text"), ("字节序", "byte_order", "text"), ("字序", "word_order", "text"), *LIMIT_COLUMNS,
]  # fmt: skip
UA_COLUMNS = [*COMMON_HEAD, ("NodeId", "node_id", "text"), *LIMIT_COLUMNS]
SHEETS = {"modbus": ("Modbus点", MODBUS_COLUMNS, "地址"), "opcua": ("OPC UA点", UA_COLUMNS, "NodeId")}
SOURCE_SHEET = "数据源"

# 类别 -> 内部编号与默认值。默认值即修复轮16 之前生产基础配置里各类点的值（终版“默认越限阈值与回差”“默认极值阈值”）。
TEMPERATURE = {"内壁温度": "temp_inner", "外壁温度": "temp_outer"}
PROCESS = {
    "机组负荷": "load_mw", "水冷壁流量": "waterwall_flow", "主汽流量": "main_steam_flow", "主汽压力": "main_steam_pressure",
    "给水温度": "feedwater_temperature", "再热温度": "reheat_temperature", "工质温度": "working_fluid_temperature",
    "炉膛负压": "furnace_pressure",
}  # fmt: skip
PROCESS_LABELS = {"水冷壁流量": "水冷壁工质流量", "主汽流量": "过热蒸汽流量", "主汽压力": "主蒸汽压力", "再热温度": "再热蒸汽温度"}
BRIDGE, OTHER = "电桥", "其他"
CATEGORIES = [*TEMPERATURE, BRIDGE, *PROCESS, OTHER]
WALL = {"unit": "degC", "high": 490, "hysteresis_high": 3, "extreme_low": -50, "extreme_high": 800}
DEFAULTS = {
    "内壁温度": WALL, "外壁温度": WALL,
    BRIDGE: {"unit": "mV/V", "extreme_low": -30, "extreme_high": 30},
    "机组负荷": {"unit": "MW", "high": 693, "hysteresis_high": 10, "extreme_low": -20, "extreme_high": 800},
    "水冷壁流量": {"unit": "t/h", "low": 925, "hysteresis_low": 90, "extreme_low": -50, "extreme_high": 3000},
    "主汽流量": {"unit": "t/h", "high": 2223, "hysteresis_high": 10, "extreme_low": -50, "extreme_high": 3000},
    "主汽压力": {"unit": "MPa", "high": 30.0, "hysteresis_high": 0.15, "extreme_low": -0.1, "extreme_high": 40},
    "给水温度": {"unit": "degC", "low": 250, "hysteresis_low": 6, "high": 320, "hysteresis_high": 1, "extreme_low": -20, "extreme_high": 400},
    "再热温度": {"unit": "degC", "high": 630, "hysteresis_high": 1, "extreme_low": -50, "extreme_high": 800},
    "工质温度": {"unit": "degC", "high": 460, "hysteresis_high": 3, "extreme_low": -50, "extreme_high": 800},
    "炉膛负压": {"unit": "kPa", "low": -0.5, "hysteresis_low": 0.05, "high": 0.5, "hysteresis_high": 0.05, "extreme_low": -5, "extreme_high": 5},
    OTHER: {},
}  # fmt: skip
DERIVED = {  # 生成的派生量的单位、阈值与极值（同上，取自原生产基础配置）
    "delta_t": {"operation": "difference", "unit": "degC", "low": -55, "hysteresis_low": 4, "high": 55, "hysteresis_high": 4,
                "extreme_low": -300, "extreme_high": 300},
    "stress": {"operation": "thermal_stress", "unit": "MPa", "extreme_low": -500, "extreme_high": 500},
    "strain": {"operation": "bridge_strain", "unit": "με", "low": -690, "hysteresis_low": 55, "high": 690, "hysteresis_high": 55,
               "extreme_low": -5000, "extreme_high": 5000},
    "strain_change": {"operation": "change", "unit": "με", "extreme_low": -10000, "extreme_high": 10000},
    "peak": {"operation": "absolute_peak", "unit": "MPa", "extreme_low": -500, "extreme_high": 500},
}  # fmt: skip
CHOICES = {
    "类别": CATEGORIES, "类型": ["modbus", "opcua"], "启用": [YES, NO], "寄存器类型": ["holding", "input"],
    "数据类型": ["uint16", "int16", "uint32", "int32", "float32", "float64"], "字节序": ["big", "little"], "字序": ["big", "little"],
}  # fmt: skip
GUIDE = [
    "每行填一个要订阅的点。不预填任何点：订阅什么就填什么。内部编号、派生量和页面测站都由点表工具生成，表里不填。",
    "类别（下拉选）：内壁温度、外壁温度、电桥、八类过程量（各最多一个）、其他。",
    "温度对号：只给内壁、外壁温度填，正整数。同一个号的一内一外是一对，生成内外壁温差与热应力估算；全部温度对再生成一个热应力峰值。",
    "电桥：每路生成换算应变、应变变化量和一个测站色块，测站名称用“显示名称”。电桥不与温度关联。",
    "显示名称：现场位号或中文名，页面上显示它；不填则按类别和编号起名。“备注”只作说明，不进配置。",
    "接入：填“地址”（Modbus 点）或“NodeId”（OPC UA 点）。这两列留空的行不接入，转换时列出，不报错。",
    "单位、上下限、回差、极值留空时按类别取默认值（转换时列出）；要去掉某个默认值，在那一格填“无”。“其他”类没有默认值，单位必填。",
    "比例、偏移留空为 1、0；“无效值”可填多个，用分号分隔。",
    "“数据源”表的“编号”要与点所在行的“数据源”一致；OPC UA 的用户名、口令只填环境变量名，不填明文。",
]


class SheetError(Exception):
    def __init__(self, problems: list[str]):
        super().__init__("\n".join(problems))
        self.problems = problems


def load_json(path: Path) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def _cell(value):
    if isinstance(value, bool):
        return YES if value else NO
    if isinstance(value, (list, tuple)):
        return ";".join(str(v) for v in value) if value else None
    return value


# ------------------------------------------------------------------------------------------------ template


def template(base: dict, out: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    Settings.model_validate(base)  # 基础配置本身先要有效
    book = Workbook()
    guide = book.active
    guide.title = "说明"
    for line in GUIDE:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="FFF2CC")

    def sheet(title, columns, rows, key_headers=()):
        ws = book.create_sheet(title)
        ws.append([header for header, _, _ in columns])
        for row in rows:
            ws.append(row)
        for k, (header, _, _) in enumerate(columns, start=1):
            cell = ws.cell(1, k)
            cell.font = bold
            if header in key_headers:
                cell.fill = fill
            letter = cell.column_letter
            ws.column_dimensions[letter].width = 30 if header in ("备注", "显示名称", "NodeId", "端点") else 12
            if header in CHOICES:
                rule = DataValidation(type="list", formula1='"' + ",".join(CHOICES[header]) + '"', allow_blank=True)
                ws.add_data_validation(rule)
                rule.add(f"{letter}2:{letter}2000")
        ws.freeze_panes = "B2"

    sheet(SOURCE_SHEET, SOURCE_COLUMNS, [[_cell(src.get(key)) for _, key, _ in SOURCE_COLUMNS] for src in base.get("sources", [])])
    for title, columns, key_header in SHEETS.values():
        sheet(title, columns, [], ("类别", "温度对号", key_header))
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)


# ------------------------------------------------------------------------------------------------- reading


def _parse(value, kind: str):
    """单元格 -> 配置值；None 表示留空，EXPLICIT_NONE 表示写了“无”。格式不对抛 ValueError（文字给填表的人看）。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, str) and value.strip() == NONE:
        return EXPLICIT_NONE
    if kind == "text":
        return str(value).strip() if not isinstance(value, float) or not value.is_integer() else str(int(value))
    if kind == "bool":
        text = str(value).strip()
        if text in (YES, "true", "True", "1"):
            return True
        if text in (NO, "false", "False", "0"):
            return False
        raise ValueError(f"应填“{YES}”或“{NO}”，现为 {value!r}")
    try:
        if kind == "int":
            number = float(str(value).strip()) if not isinstance(value, (int, float)) else value
            if isinstance(number, float) and not number.is_integer():
                raise ValueError
            return int(number)
        if kind == "float":
            return float(str(value).strip()) if not isinstance(value, (int, float)) else float(value)
        if kind == "floats":
            parts = str(value).replace("；", ";").replace(",", ";").split(";")
            return [float(p) for p in parts if p.strip()]
    except ValueError:
        what = {"int": "整数", "float": "数值", "floats": "数值（多个用分号分隔）"}[kind]
        raise ValueError(f"应为{what}，现为 {value!r}") from None
    raise AssertionError(kind)


def _rows(book, title: str, columns):
    """(行号, {表头: 单元格}) for each non-empty row; missing headers are an error."""
    if title not in book.sheetnames:
        raise SheetError([f"缺少工作表“{title}”"])
    ws = book[title]
    headers = [c.value for c in ws[1]]
    missing = [h for h, _, _ in columns if h not in headers]
    if missing:
        raise SheetError([f"{title}：缺少列 {'、'.join(missing)}（这张表是旧版点表？用“生成点表模板”出一份新的）"])
    at = {h: headers.index(h) for h, _, _ in columns}
    for number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        cells = {h: (row[i] if i < len(row) else None) for h, i in at.items()}
        if any(v is not None and str(v).strip() for v in cells.values()):
            yield number, cells


def _fields(title, number, cells, columns, problems):
    out = {}
    for header, key, kind in columns:
        if key is None:
            continue
        try:
            value = _parse(cells[header], kind)
        except ValueError as error:
            problems.append(f"{title} 第{number}行 {header}：{error}")
            continue
        if value is not None:
            out[key] = value
    return out


def _model_problems(error: ValidationError, where, columns, name):
    headers = {key: header for header, key, _ in columns if key}
    found = []
    for item in error.errors():
        field = item["loc"][0] if item["loc"] else None
        column = headers.get(field, "") if field else ""
        message = str(item.get("msg", "")).removeprefix("Value error, ")
        found.append(f"{where} {name}{('：' + column) if column else ''}：{message}")
    return found


def read_sheet(path: Path):
    """(数据源, 行) ；行是 dict：where、kind、category、pair、shown、raw（点的字段，未接入时为 None）。"""
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True)
    problems, sources, rows = [], [], []
    for number, cells in _rows(book, SOURCE_SHEET, SOURCE_COLUMNS):
        raw = _fields(SOURCE_SHEET, number, cells, SOURCE_COLUMNS, problems)
        try:
            sources.append(Source.model_validate(raw).model_dump(exclude_defaults=True) | {"id": raw.get("id"), "kind": raw.get("kind")})
        except ValidationError as error:
            problems += _model_problems(error, f"{SOURCE_SHEET} 第{number}行", SOURCE_COLUMNS, raw.get("id", ""))
    kinds = {s["id"]: s["kind"] for s in sources}
    for kind, (title, columns, key_header) in SHEETS.items():
        for number, cells in _rows(book, title, columns):
            where = f"{title} 第{number}行"
            raw = _fields(title, number, cells, columns, problems)
            category = raw.pop("category", None)
            pair = raw.pop("pair", None)
            if category not in CATEGORIES:
                problems.append(f"{where} 类别：{'未填' if category is None else f'不认识“{category}”'}，应从下拉里选（{'、'.join(CATEGORIES)}）")
                continue
            if category in TEMPERATURE:
                if not isinstance(pair, int) or pair < 1:
                    problems.append(f"{where} 温度对号：{category}要填温度对号（正整数），同一个号的一内一外是一对")
                    continue
            elif pair is not None:
                problems.append(f"{where} 温度对号：只给内壁、外壁温度填，{category}不填")
                continue
            connected = _parse(cells[key_header], "text") not in (None, EXPLICIT_NONE)
            if connected and raw.get("source") not in kinds:
                problems.append(f"{where} 数据源：{raw.get('source')!r} 不在“{SOURCE_SHEET}”表里")
                continue
            if connected and kinds[raw["source"]] != kind:
                problems.append(f"{where} 数据源：{raw['source']} 是 {kinds[raw['source']]}，不能放在“{title}”表")
                continue
            rows.append({"where": where, "kind": kind, "columns": columns, "category": category, "pair": pair,
                         "shown": _parse(cells["显示名称"], "text"), "raw": raw if connected else None})  # fmt: skip
    if problems:
        raise SheetError(problems)
    return sources, rows


# ------------------------------------------------------------------------------------------------- convert


def assign(rows: list[dict]) -> tuple[list[dict], list[str]]:
    """Every row its internal id; problems for a doubled temperature side or process quantity."""
    problems, taken, bridges, others = [], {}, 0, 0
    for row in rows:
        category = row["category"]
        if category in TEMPERATURE:
            row["id"] = f"{TEMPERATURE[category]}_{row['pair']:02d}"
        elif category == BRIDGE:
            bridges += 1
            row["id"] = f"bridge_mvv_{bridges:02d}"
        elif category in PROCESS:
            row["id"] = PROCESS[category]
        else:
            others += 1
            row["id"] = f"other_{others:02d}"
        if row["id"] in taken:
            what = f"温度对 {row['pair']:02d} 的{category}" if category in TEMPERATURE else category
            problems.append(f"{row['where']}：{what}已经在 {taken[row['id']]} 填过，每个只能有一个")
        else:
            taken[row["id"]] = row["where"]
    return rows, problems


def point_of(row: dict, used_defaults: dict, problems: list[str]) -> dict | None:
    raw, defaults = dict(row["raw"]), DEFAULTS[row["category"]]
    filled = []
    for key, value in defaults.items():
        side = key.removeprefix("hysteresis_")
        if key != side and raw.get(side) is EXPLICIT_NONE:
            continue  # a limit taken away with “无” takes its default hysteresis with it
        if key not in raw:
            raw[key] = value
            filled.append(key)
    for key in [k for k, v in raw.items() if v is EXPLICIT_NONE]:
        del raw[key]
    if filled:
        used_defaults.setdefault(row["category"], set()).update(filled)
    raw["id"] = row["id"]
    try:
        point = Point.model_validate(raw)
    except ValidationError as error:
        problems += _model_problems(error, row["where"], row["columns"], row["shown"] or row["category"])
        return None
    if row["kind"] == "opcua":
        try:
            check_node_id(point)
        except ValueError as error:
            problems.append(f"{row['where']} NodeId：{str(error).split(': ', 1)[-1]}")
            return None
    return point.model_dump(exclude_defaults=True) | {"id": point.id, "source": point.source, "unit": point.unit}


DEFAULT_NAMES = {"unit": "单位", "low": "下限", "hysteresis_low": "下限回差", "high": "上限", "hysteresis_high": "上限回差",
                 "extreme_low": "极小值", "extreme_high": "极大值"}  # fmt: skip


def convert(base: dict, sheet: Path) -> tuple[dict, list[str]]:
    """(新配置, 报告行)。点表有错时抛 SheetError，不产生配置。"""
    sources, rows = read_sheet(sheet)
    rows, problems = assign(rows)
    connected = [r for r in rows if r["raw"] is not None]
    used_defaults: dict[str, set] = {}
    points = {}
    for row in connected:
        point = point_of(row, used_defaults, problems)
        if point is not None:
            points[row["id"]] = point
    if problems:
        raise SheetError(problems)
    kinds = {s["id"]: s["kind"] for s in sources}
    report = [f"接入 {len(points)} 个点（Modbus {sum(kinds[p['source']] == 'modbus' for p in points.values())}，"
              f"OPC UA {sum(kinds[p['source']] == 'opcua' for p in points.values())}）；未接入 {len(rows) - len(connected)} 个"]  # fmt: skip
    if len(rows) > len(connected):
        report.append("未接入（地址或 NodeId 留空）：" + "、".join(f"{r['where']}（{r['shown'] or r['category']}）" for r in rows if r["raw"] is None))
    for category in CATEGORIES:
        if category in used_defaults:
            report.append(f"{category}按类别默认值补了：" + "、".join(DEFAULT_NAMES[k] for k in DEFAULT_NAMES if k in used_defaults[category]))
    shown = {r["id"]: r["shown"] for r in connected}

    # the points in a fixed order: bridges, temperature pairs, process quantities, others
    pairs = sorted({r["pair"] for r in connected if r["category"] in TEMPERATURE})
    bridges = sorted(i for i in points if i.startswith("bridge_mvv_"))
    order = [*bridges, *(f"{side}_{n:02d}" for n in pairs for side in ("temp_inner", "temp_outer")), *PROCESS.values(),
             *sorted(i for i in points if i.startswith("other_"))]  # fmt: skip
    labels = {}
    for category, prefix in TEMPERATURE.items():
        for n in pairs:
            pid = f"{prefix}_{n:02d}"
            if pid in points:
                labels[pid] = shown[pid] or f"温度对{n:02d} {category}"
    for k, pid in enumerate(bridges, start=1):
        labels[pid] = shown[pid] or f"电桥{k:02d}"
    for category, pid in PROCESS.items():
        if pid in points:
            labels[pid] = shown[pid] or PROCESS_LABELS.get(category, category)
    for pid in points:
        if pid.startswith("other_"):
            labels[pid] = shown[pid] or f"其他{pid[-2:]}"

    derived, stations = [], []
    model_id = next((m["id"] for m in base.get("models", [])), None)
    complete = [n for n in pairs if f"temp_inner_{n:02d}" in points and f"temp_outer_{n:02d}" in points]
    for n in pairs:
        if n not in complete:
            side = "内壁" if f"temp_inner_{n:02d}" in points else "外壁"
            report.append(f"温度对{n:02d} 只接入了{side}：照常接入，不生成温差与热应力")
    for n in complete:
        inputs = [f"temp_inner_{n:02d}", f"temp_outer_{n:02d}"]
        derived.append({"id": f"delta_t_{n:02d}", "inputs": inputs, **DERIVED["delta_t"]})
        labels[f"delta_t_{n:02d}"] = f"温度对{n:02d} 内外壁温差"
        if model_id is not None:
            derived.append({"id": f"stress_{n:02d}", "inputs": inputs, **DERIVED["stress"], "model_id": model_id})
            labels[f"stress_{n:02d}"] = f"温度对{n:02d} 热应力估算"
    columns = 2
    rows_of_blocks = max(1, -(-len(bridges) // columns))
    for k, pid in enumerate(bridges, start=1):
        n = pid[-2:]
        derived.append({"id": f"strain_{n}", "inputs": [pid], **DERIVED["strain"]})
        derived.append({"id": f"strain_change_{n}", "inputs": [f"strain_{n}"], **DERIVED["strain_change"]})
        labels[f"strain_{n}"], labels[f"strain_change_{n}"] = f"{labels[pid]} 换算应变", f"{labels[pid]} 应变变化量"
        row_index = (k - 1) // columns
        stations.append({"id": f"S{n}", "label": labels[pid], "x": 25 if (k - 1) % columns == 0 else 75,
                         "y": round(15 + row_index * (70 / max(rows_of_blocks - 1, 1)), 1),
                         "strain": f"strain_{n}", "strain_change": f"strain_change_{n}"})  # fmt: skip
    stresses = [d["id"] for d in derived if d["operation"] == "thermal_stress"]
    if stresses:
        derived.append({"id": "thermal_stress_peak", "inputs": stresses, **DERIVED["peak"]})
        labels["thermal_stress_peak"] = "热应力代理绝对峰值（保留符号）"
    report.append(f"生成：温差与热应力 {len(complete)} 对，换算应变与应变变化量 {len(bridges)} 路，"
                  f"测站 {len(stations)} 个，热应力峰值 {1 if stresses else 0} 个")  # fmt: skip

    config = json.loads(json.dumps(base))  # 基础配置的其余部分原样保留
    used = {p["source"] for p in points.values()}
    for source in sources:
        if source.get("enabled") and source["id"] not in used:
            source["enabled"] = False
            report.append(f"数据源 {source['id']} 没有接入的点：改为不启用")
    config["sources"] = sources
    config["points"] = [points[i] for i in order if i in points]
    config["derived"] = derived
    ui = config.setdefault("ui", {})
    ui["stations"], ui["point_labels"] = stations, labels
    ui["process_points"] = [pid for pid in PROCESS.values() if pid in points]
    simulation = config.get("simulation")
    if simulation and simulation.get("enabled"):  # 模拟现场演练：模拟器给每个点一个信号
        signals = {s["point"]: s for s in simulation.get("signals", [])}
        constant = [p["id"] for p in config["points"] if p["id"] not in signals]
        simulation["signals"] = [signals.get(p["id"]) or {"point": p["id"], "base": (p["extreme_low"] + p["extreme_high"]) / 2
                                 if "extreme_low" in p and "extreme_high" in p else 0.0} for p in config["points"]]  # fmt: skip
        if constant:
            report.append("模拟现场为基础配置里没有的点给恒定值（极值中点，无极值时为0）：" + "、".join(constant))
        for scenario in simulation.get("scenarios", []):
            scenario["points"] = [x for x in scenario.get("points", []) if x in points]
            scenario["offsets"] = {k: v for k, v in scenario.get("offsets", {}).items() if k in points}
    try:
        Settings.model_validate(config)
    except ValidationError as error:
        raise SheetError(["生成的配置没有通过后端的配置校验：" + "; ".join(
            str(e.get("msg", "")).removeprefix("Value error, ") for e in error.errors()
        )]) from None  # fmt: skip
    return config, report


def main(argv=None):
    parser = argparse.ArgumentParser(description="点表 Excel 与部署配置 JSON 的互转")
    sub = parser.add_subparsers(dest="command", required=True)
    t = sub.add_parser("template", help="生成空白点表")
    t.add_argument("--base", type=Path, required=True)
    t.add_argument("--out", type=Path, required=True)
    c = sub.add_parser("convert", help="由填好的点表生成部署配置")
    c.add_argument("--base", type=Path, required=True)
    c.add_argument("--sheet", type=Path, required=True)
    c.add_argument("--out", type=Path, required=True)
    args = parser.parse_args(argv)
    base = load_json(args.base)
    if args.command == "template":
        template(base, args.out)
        print(f"空白点表已写入 {args.out}")
        return 0
    try:
        config, report = convert(base, args.sheet)
    except SheetError as error:
        print("点表有错，没有生成配置：", file=sys.stderr)
        for line in error.problems:
            print("  " + line, file=sys.stderr)
        return 1
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(config, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for line in report:
        print(line)
    print(f"配置已写入 {args.out}（已按后端配置规则校验通过）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
