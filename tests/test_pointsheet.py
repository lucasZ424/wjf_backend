"""tools/pointsheet: the site's points table in Excel, turned into the deployment configuration (修复轮16, user 2026-11-01).

Why it matters: the site engineer knows what each point is and where to subscribe it, not our internal naming.  The table
therefore carries site facts only - a category, the pair number of a wall temperature, a tag, an address - and the tool
gives every point its internal id and builds the temperature differences, the stress proxies, the strains, the stations
and the peak from them.  Nothing is pre-filled: a point that is not in the table is not in the configuration.  A table with
a mistake produces no configuration at all, and what reaches the backend has passed the same rules as `htsg check`.
"""

import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from htsg.config import Settings

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools" / "pointsheet"))
import pointsheet  # noqa: E402

BASE = pointsheet.load_json(ROOT / "config/production.base.json")
RESEARCH = pointsheet.load_json(ROOT / "config/research-sim.json")


def blank(tmp_path, base=BASE):
    path = tmp_path / "点表.xlsx"
    pointsheet.template(base, path)
    book = load_workbook(path)
    sources = book["数据源"]
    heads = [c.value for c in sources[1]]
    for row in sources.iter_rows(min_row=2):
        cells = dict(zip(heads, row))
        cells["启用"].value = "是"
        if cells["类型"].value == "modbus":
            cells["主机"].value, cells["端口"].value = "127.0.0.1" if base is RESEARCH else "192.168.1.10", 15020 if base is RESEARCH else 502
        else:
            cells["端点"].value = "opc.tcp://127.0.0.1:14840/htsg/" if base is RESEARCH else "opc.tcp://192.168.1.20:4840/"
    book.save(path)
    return path


def add(path, title, **values):
    book = load_workbook(path)
    ws = book[title]
    heads = [c.value for c in ws[1]]
    ws.append([values.get(h) for h in heads])
    book.save(path)


def bridge(path, n, **extra):
    add(path, "Modbus点", 类别="电桥", 显示名称=f"SG-{n:02d}", 数据源="mb_field", 从站号=1, 地址=2 * (n - 1), 数据类型="float32", **extra)


def wall(path, side, pair, **extra):
    tag = f"TE-{pair:02d}{'A' if side == '内壁温度' else 'B'}"
    values = {"类别": side, "温度对号": pair, "显示名称": tag, "数据源": "ua_field", "NodeId": f"ns=2;s={tag}", **extra}
    add(path, "OPC UA点", **values)


def site(tmp_path, pairs=(1, 2), bridges=2):
    path = blank(tmp_path)
    for n in range(1, bridges + 1):
        bridge(path, n)
    for pair in pairs:
        wall(path, "内壁温度", pair)
        wall(path, "外壁温度", pair)
    add(path, "OPC UA点", 类别="机组负荷", 显示名称="LOAD", 数据源="ua_field", NodeId="ns=2;s=LOAD")
    return path


def test_the_blank_table_carries_no_point_and_offers_the_categories(tmp_path):
    """用户 11-01: no default points.  The sources are named, the point sheets are only headers and choices."""
    assert BASE["points"] == [] and BASE["derived"] == [] and BASE["ui"]["stations"] == [], "the base has no default points"
    book = load_workbook(blank(tmp_path))
    for title in ("Modbus点", "OPC UA点"):
        ws = book[title]
        heads = [c.value for c in ws[1]]
        assert ws.max_row == 1, f"{title} has no pre-filled row"
        assert heads[:2] == ["类别", "温度对号"] and "内部编号" not in heads, "the internal id is not the site's business"
        choices = [v.formula1 for v in ws.data_validations.dataValidation if "内壁温度" in (v.formula1 or "")]
        assert choices and all(c in choices[0] for c in pointsheet.CATEGORIES)
    assert [c.value for c in book["数据源"]["A"][1:]] == ["mb_field", "ua_field"]


def test_ids_derived_values_stations_and_labels_come_from_the_site_facts(tmp_path):
    config, report = pointsheet.convert(BASE, site(tmp_path, pairs=(3, 7)))
    settings = Settings.model_validate(config)
    assert [p.id for p in settings.points] == ["bridge_mvv_01", "bridge_mvv_02", "temp_inner_03", "temp_outer_03",
                                               "temp_inner_07", "temp_outer_07", "load_mw"]  # fmt: skip
    derived = {d.id: d for d in settings.derived}
    assert derived["delta_t_07"].inputs == ("temp_inner_07", "temp_outer_07"), "inner minus outer, as the pages read it"
    assert derived["stress_03"].operation == "thermal_stress" and derived["stress_03"].model_id == settings.models[0].id
    assert derived["strain_change_02"].inputs == ("strain_02",) and derived["strain_02"].inputs == ("bridge_mvv_02",)
    assert derived["thermal_stress_peak"].operation == "absolute_peak"
    assert derived["thermal_stress_peak"].inputs == ("stress_03", "stress_07")
    stations = settings.ui.stations
    assert [(s.id, s.label, s.strain) for s in stations] == [("S01", "SG-01", "strain_01"), ("S02", "SG-02", "strain_02")]
    assert all(s.temperature is None and s.delta_temperature is None for s in stations), "a bridge is not tied to a pair"
    labels = settings.ui.point_labels
    assert labels["temp_inner_07"] == "TE-07A" and labels["delta_t_07"] == "温度对07 内外壁温差"
    assert settings.ui.process_points == ("load_mw",)
    assert report[-1].startswith("生成：温差与热应力 2 对，换算应变与应变变化量 2 路，测站 2 个，热应力峰值 1 个")


def test_an_empty_cell_takes_the_category_default_and_none_takes_it_away(tmp_path):
    """用户 11-01: a threshold left empty is the category's default; “无” removes it, with its hysteresis."""
    path = site(tmp_path, pairs=(1,))
    wall(path, "内壁温度", 2, 上限="无")
    wall(path, "外壁温度", 2, 上限=500)
    add(path, "OPC UA点", 类别="其他", 显示名称="汽包水位", 数据源="ua_field", 单位="mm", NodeId="ns=2;s=LT101")
    config, report = pointsheet.convert(BASE, path)
    points = {p["id"]: p for p in config["points"]}
    assert points["temp_inner_01"]["high"] == 490 and points["temp_inner_01"]["unit"] == "degC"
    assert "high" not in points["temp_inner_02"] and "hysteresis_high" not in points["temp_inner_02"]
    assert points["temp_outer_02"]["high"] == 500 and points["temp_outer_02"]["hysteresis_high"] == 3
    assert points["other_01"] == {"id": "other_01", "source": "ua_field", "unit": "mm", "node_id": "ns=2;s=LT101"}, "no default"
    assert any(line.startswith("内壁温度按类别默认值补了：单位、上限") for line in report)


def test_a_row_without_address_is_not_connected_and_a_half_pair_gets_no_derived_value(tmp_path):
    path = site(tmp_path, pairs=(1,))
    wall(path, "内壁温度", 2)
    wall(path, "外壁温度", 2, NodeId=None)
    config, report = pointsheet.convert(BASE, path)
    ids = {p["id"] for p in config["points"]} | {d["id"] for d in config["derived"]}
    assert "temp_inner_02" in ids and "temp_outer_02" not in ids and "delta_t_02" not in ids and "stress_02" not in ids
    assert any("OPC UA点 第" in line and "TE-02B" in line for line in report), "what is not connected is named"
    assert "温度对02 只接入了内壁：照常接入，不生成温差与热应力" in report


@pytest.mark.parametrize(
    "rows, message",
    [
        ([("OPC UA点", {"类别": "壁温", "数据源": "ua_field", "NodeId": "ns=2;s=x"})], "类别：不认识“壁温”"),
        ([("OPC UA点", {"类别": "内壁温度", "数据源": "ua_field", "NodeId": "ns=2;s=x"})], "温度对号：内壁温度要填温度对号"),
        ([("OPC UA点", {"类别": "机组负荷", "温度对号": 1, "数据源": "ua_field", "NodeId": "ns=2;s=x"})], "只给内壁、外壁温度填"),
        ([("OPC UA点", {"类别": "外壁温度", "温度对号": 1, "数据源": "ua_field", "NodeId": "ns=2;s=x"})], "温度对 01 的外壁温度已经在"),
        ([("OPC UA点", {"类别": "机组负荷", "数据源": "ua_field", "NodeId": "ns=2;s=y"})], "机组负荷已经在"),
        ([("OPC UA点", {"类别": "其他", "数据源": "ua_field", "NodeId": "ns=2;s=z"})], "其他"),
        ([("OPC UA点", {"类别": "工质温度", "数据源": "nowhere", "NodeId": "ns=2;s=w"})], "数据源：'nowhere' 不在"),
    ],
)
def test_a_mistake_names_the_sheet_and_row_and_produces_nothing(tmp_path, rows, message):
    path = site(tmp_path, pairs=(1,))
    for title, values in rows:
        add(path, title, **values)
    with pytest.raises(pointsheet.SheetError) as error:
        pointsheet.convert(BASE, path)
    assert any("OPC UA点 第" in line and message in line for line in error.value.problems), error.value.problems


def test_an_old_table_with_internal_ids_is_refused_with_a_hint(tmp_path):
    path = blank(tmp_path)
    book = load_workbook(path)
    book["OPC UA点"].cell(1, 1).value = "内部编号"
    book.save(path)
    with pytest.raises(pointsheet.SheetError, match="旧版点表"):
        pointsheet.convert(BASE, path)


def test_an_unused_source_is_switched_off_and_named(tmp_path):
    path = blank(tmp_path)
    wall(path, "内壁温度", 1)
    wall(path, "外壁温度", 1)
    config, report = pointsheet.convert(BASE, path)
    assert {s["id"]: s["enabled"] for s in config["sources"]} == {"mb_field": False, "ua_field": True}
    assert "数据源 mb_field 没有接入的点：改为不启用" in report
    assert Settings.model_validate(config).ui.stations == ()


def test_a_research_base_gives_the_simulated_field_a_signal_for_every_point(tmp_path):
    """The same table drives a simulated site: points the research configuration knows keep their signals, the rest send
    a constant, and the simulator's plant model recognises the ids the tool gives (temp_inner_NN, bridge_mvv_NN ...)."""
    path = blank(tmp_path, RESEARCH)
    bridge(path, 1)
    wall(path, "内壁温度", 1)
    wall(path, "外壁温度", 1)
    add(path, "OPC UA点", 类别="其他", 显示名称="汽包水位", 数据源="ua_field", 单位="mm", NodeId="ns=2;s=drum")
    config, report = pointsheet.convert(RESEARCH, path)
    settings = Settings.model_validate(config)
    signals = {s.point: s for s in settings.simulation.signals}
    assert set(signals) == {p.id for p in settings.points}
    known = next(s for s in RESEARCH["simulation"]["signals"] if s["point"] == "temp_inner_01")
    assert (signals["temp_inner_01"].base, signals["temp_inner_01"].amplitude) == (known["base"], known["amplitude"])
    assert signals["other_01"].base == 0.0 and any("other_01" in line for line in report)
